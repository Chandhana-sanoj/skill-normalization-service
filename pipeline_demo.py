
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from app.repositories.canonical_skill_repository import (
    CanonicalSkillRepository,
)
from app.services.normalizer import SkillNormalizer
from app.services.vector_search import VectorSearcher
from app.services.multi_skill_extractor import MultiSkillExtractor


EXCEL_FILE = Path("pipeline_results.xlsx")
ORGANIZATION_ID = 1

RESULT_COLUMNS = [
    "Input",
    "Normalized",
    "Version",
    "Description",
    "Category",
    "Status",
]


# --------------------------------------------------
# 1. Save results to Excel
# --------------------------------------------------

def save_results_to_excel(results):
    """Append results to Excel without deleting previous history."""

    if not results:
        return

    if EXCEL_FILE.exists():
        workbook = load_workbook(EXCEL_FILE)
        sheet = workbook.active

        # Add headers if the existing workbook is empty.
        if sheet.max_row == 1 and sheet.cell(1, 1).value is None:
            sheet.append(RESULT_COLUMNS)

            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    "solid",
                    fgColor="1F4E78",
                )
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Normalization Results"
        sheet.append(RESULT_COLUMNS)

        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                "solid",
                fgColor="1F4E78",
            )

    for result in results:
        sheet.append([
            result.get("input"),
            result.get("skill"),
            result.get("version"),
            result.get("description"),
            result.get("category"),
            result.get("status"),
        ])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column in sheet.columns:
        max_length = max(
            len(str(cell.value or ""))
            for cell in column
        )

        sheet.column_dimensions[
            column[0].column_letter
        ].width = min(max_length + 2, 45)

    workbook.save(EXCEL_FILE)


# --------------------------------------------------
# 2. Load catalogue from PostgreSQL
# --------------------------------------------------

catalogue_repository = CanonicalSkillRepository()

(
    alias_map,
    canonical_map,
    description_map,
    category_map,
) = catalogue_repository.load_catalogue()


# --------------------------------------------------
# 3. Initialize services
# --------------------------------------------------

multi_skill_extractor = MultiSkillExtractor(
    alias_map=alias_map,
    canonical_map=canonical_map,
)

vector_searcher = VectorSearcher()

normalizer = SkillNormalizer(
    alias_map=alias_map,
    canonical_map=canonical_map,
    vector_searcher=vector_searcher,
    organization_id=ORGANIZATION_ID,
    description_map=description_map,
    category_map=category_map,
)


# --------------------------------------------------
# 4. Process terminal input
# --------------------------------------------------

def process_input(user_input):
    """Extract skills, normalize them, and save results."""

    if not user_input.strip():
        print("Empty input. Please enter a skill or phrase.")
        return

    try:
        # Extract detailed candidates so versions remain separate.
        candidates = multi_skill_extractor.extract_detailed(
            user_input
        )

        if not candidates:
            print("No skills extracted.")
            return

        results = []

        for candidate in candidates:
            # Pass the skill text and version separately.
            result = normalizer.normalize(
                candidate.text,
                version=candidate.version,
            )

            excel_result = {
                "input": user_input,
                "skill": result.get("skill"),
                "version": result.get("version"),
                "description": result.get("description"),
                "category": result.get("category"),
                "status": result.get("status"),
            }

            results.append(excel_result)

        save_results_to_excel(results)

        print(f"Processed: {user_input}")
        print(f"Skills found: {len(results)}")
        print(f"Saved to: {EXCEL_FILE.resolve()}")

        # Display the results for immediate verification.
        for result in results:
            print(
                f"  Skill: {result['skill']} | "
                f"Version: {result['version']} | "
                f"Status: {result['status']}"
            )

    except ValueError as exc:
        print(f"Normalization error: {exc}")


# --------------------------------------------------
# 5. Interactive terminal loop
# --------------------------------------------------

def main():
    print("Skill Normalization Pipeline")
    print("Enter a skill or phrase.")
    print("Type 'exit' to stop.")
    print("-" * 40)

    while True:
        try:
            user_input = input("\nEnter skill: ").strip()

            if user_input.lower() == "exit":
                break

            process_input(user_input)

        except KeyboardInterrupt:
            print("\nStopping pipeline.")
            break

    print("\nPipeline stopped.")
    print(f"Results file: {EXCEL_FILE.resolve()}")


if __name__ == "__main__":
    main()
